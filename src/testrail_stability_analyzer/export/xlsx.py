"""Native XLSX stability report (Excel, Google Sheets), written with the optional ``openpyxl`` package.

Worksheets, in order: ``Stability Overview`` (one compact status cell per
application/platform, with a short note), ``Stability Details`` (every metric,
as numeric percentage and date cells), and ``Report Info``. The content is the
same as the XML report's (see :mod:`.report_data`).

``openpyxl`` is imported only when a workbook is rendered, so collection and
XML export work without it; :func:`xlsx_available` lets callers check first.
Every string is passed through :func:`.report_data.safe_text` and stored as a
text cell: no formulas, hyperlinks, or external links are written.
"""

from __future__ import annotations

import importlib.util
import io
from collections.abc import Mapping, Sequence
from datetime import UTC, datetime
from typing import Any, Final

from ..analysis.models import AnalysisReport
from ..collection.models import CollectionReport
from ..config.settings import Settings
from .report_data import GREEN, NEUTRAL, RED, YELLOW, Cell, ReportContent, Worksheet, build_report, safe_text

INSTALL_HINT: Final = (
    'XLSX reports need the optional openpyxl package: pip install "openpyxl>=3.1,<4" '
    "(or: poetry install --extras xlsx)"
)
AUTHOR: Final = "testrail-stability-analyzer"
# (font colour, fill colour) per report style; the same palette as the XML report.
_COLOURS: Final = {
    GREEN: ("006100", "C6EFCE"),
    YELLOW: ("9C5700", "FFEB9C"),
    RED: ("9C0006", "FFC7CE"),
    NEUTRAL: ("3F3F3F", "EDEDED"),
}
_HEADER_FILL: Final = "D9D9D9"
_NUMBER_FORMATS: Final = {"percent": "0.0%", "datetime": "yyyy-mm-dd hh:mm"}
_POINTS_PER_CHARACTER: Final = 6.5  # converts the shared column widths (points) to Excel character units


class XlsxUnavailableError(RuntimeError):
    """``openpyxl`` is not installed."""


def xlsx_available() -> bool:
    try:
        return importlib.util.find_spec("openpyxl") is not None
    except (ImportError, ValueError):  # e.g. blocked through sys.modules
        return False


def render_xlsx(
    collection: CollectionReport,
    analysis: AnalysisReport,
    settings: Settings,
    *,
    case_types: Mapping[int, str] | None = None,
    notes: Sequence[str] = (),
    content: ReportContent | None = None,
) -> bytes:
    """Render the workbook as XLSX bytes; an already built ``content`` is reused."""
    try:
        from openpyxl import Workbook
    except ImportError:
        raise XlsxUnavailableError(INSTALL_HINT) from None
    content = content or build_report(collection, analysis, settings, case_types=case_types, notes=notes)
    workbook = Workbook()
    workbook.remove(workbook.active)
    styles = _Styles()
    for sheet, frozen in ((content.overview, "B2"), (content.details, "B2"), (content.info, "A2")):
        _write_sheet(workbook.create_sheet(sheet.name), sheet, styles, frozen)
    collected = collection.collected_at.astimezone(UTC).replace(tzinfo=None)
    workbook.properties.creator = AUTHOR
    workbook.properties.created = collected
    workbook.properties.modified = collected
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


class _Styles:
    """openpyxl style objects, created once per workbook."""

    def __init__(self) -> None:
        from openpyxl.styles import Alignment, Font, PatternFill

        self.top = Alignment(vertical="top")
        self.wrapped = Alignment(vertical="top", wrap_text=True)
        self.header_font = Font(bold=True)
        self.header_fill = PatternFill("solid", fgColor=_HEADER_FILL)
        self.fonts = {style: Font(color=font) for style, (font, _) in _COLOURS.items()}
        self.fills = {style: PatternFill("solid", fgColor=fill) for style, (_, fill) in _COLOURS.items()}


def _write_sheet(ws: Any, sheet: Worksheet, styles: _Styles, frozen: str) -> None:
    from openpyxl.comments import Comment
    from openpyxl.utils import get_column_letter

    wrapped = {index for index, (header, _) in enumerate(sheet.columns, start=1) if header in sheet.wrap}
    for index, (header, width) in enumerate(sheet.columns, start=1):
        ws.column_dimensions[get_column_letter(index)].width = round(width / _POINTS_PER_CHARACTER, 1)
        cell = ws.cell(row=1, column=index)
        _set_text(cell, header)
        cell.font, cell.fill, cell.alignment = styles.header_font, styles.header_fill, styles.wrapped
    for row_index, row in enumerate(sheet.rows, start=2):
        for column_index, value in enumerate(row, start=1):
            cell = ws.cell(row=row_index, column=column_index)
            _set_value(cell, value)
            cell.alignment = styles.wrapped if column_index in wrapped else styles.top
            if value.style in _NUMBER_FORMATS:
                cell.number_format = _NUMBER_FORMATS[value.style]
            elif value.style in styles.fills:
                cell.font, cell.fill = styles.fonts[value.style], styles.fills[value.style]
            if value.comment:
                comment = Comment(safe_text(value.comment), AUTHOR)
                comment.width, comment.height = 260, 130
                cell.comment = comment
    ws.freeze_panes = frozen
    ws.auto_filter.ref = f"A1:{get_column_letter(len(sheet.columns))}{len(sheet.rows) + 1}"


def _set_value(cell: Any, value: Cell) -> None:
    data = value.value
    if data is None:
        return
    if isinstance(data, datetime):
        cell.value = data.astimezone(UTC).replace(tzinfo=None)  # Excel dates carry no time zone; reports use UTC
    elif isinstance(data, bool):  # never exported as a number
        _set_text(cell, "yes" if data else "no")
    elif isinstance(data, int | float):
        cell.value = data
    else:
        _set_text(cell, data)


def _set_text(cell: Any, text: str) -> None:
    """A text cell, never a formula: openpyxl would otherwise treat a leading '=' as one."""
    cell.value = safe_text(text)
    cell.data_type = "s"
