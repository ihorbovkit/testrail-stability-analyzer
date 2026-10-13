"""Blocked reason: latest-status selection in analysis and report output (XML and XLSX). Synthetic data only."""

import io
import unittest
from datetime import timedelta

from testrail_stability_analyzer.analysis import analyze
from testrail_stability_analyzer.collection import DiscoveredRun, RunClassification, RunDecision, build_outcome
from testrail_stability_analyzer.export import build_report, render_spreadsheet, render_xlsx
from testrail_stability_analyzer.models import Result, Run, Test

from . import analysis_fixtures as af
from . import fake_testrail as fx
from .analysis_fixtures import APP_A, PLAT_X
from .fake_testrail import BLOCKED, FAILED, NOW, PASSED
from .test_export_spreadsheet import table

try:
    import openpyxl
except ImportError:  # pragma: no cover - exercised only without the optional dependency
    openpyxl = None

RAW = (
    "Automation run 77\nBLOCKED REASON: Payment sandbox unavailable\nsessionId=SECRET-SESSION-42\n\n"
    "Full log: https://ci.example.invalid/job/77 ![shot](index.php?/attachments/get/9)\n"
    "Traceback (most recent call last):\n  File \"/srv/t.py\", line 3\nRAW-COMMENT-TAIL"
)
REASON = "Payment sandbox unavailable sessionId=[redacted]"
COLUMN = "App A | Platform X: Blocked Reason"


def record(result_id, status_id, days, comment=None):
    """``status_id`` ``None`` makes a comment-only record."""
    return (result_id, status_id, days, comment)


def run_outcome(run_id, *records, case_id=100, complete=True):
    test_id = run_id * 1000 + case_id
    source = DiscoveredRun(Run(id=run_id, name=f"Run {run_id}"), RunClassification(APP_A, PLAT_X), RunDecision.INCLUDED)
    results = [
        Result(id=rid, test_id=test_id, status_id=status, created_on=NOW - timedelta(days=days), comment=comment)
        for rid, status, days, comment in records
    ]
    test = Test(id=test_id, case_id=case_id, run_id=run_id, title=f"Case {case_id}")
    return build_outcome(source, test, results, af.RESOLVER, results_complete=complete)


def settings(**blocked_reason):
    return fx.settings(report={"blocked_reason": blocked_reason}) if blocked_reason else fx.settings()


def analysis_of(outcomes, config=None):
    collected = af.report(outcomes)
    config = config or settings()
    return collected, analyze(collected, config), config


def key_analysis(outcomes, config=None, case_id=100):
    _, analysis, _ = analysis_of(outcomes, config)
    return next(a for a in analysis.analyses if a.key.case_id == case_id and a.key.platform == PLAT_X)


class LatestStatusSelectionTests(unittest.TestCase):
    def test_reason_of_the_latest_verified_blocked_final(self):
        a = key_analysis([run_outcome(500, record(1, BLOCKED, 1, RAW))])
        self.assertEqual(a.blocked_reason, REASON)
        self.assertEqual(a.metrics.latest_category.value, "blocked")

    def test_triage_note_after_the_final_is_used_when_the_final_has_none(self):
        triage = record(2, None, 1, "SKIPPED REASON: Device lab closed")
        outcome = run_outcome(500, record(1, BLOCKED, 2, "no section"), triage)
        self.assertEqual(key_analysis([outcome]).blocked_reason, "Device lab closed")

    def test_final_comment_precedes_later_triage_notes(self):
        outcome = run_outcome(
            500, record(1, BLOCKED, 2, "BLOCKED REASON: Final reason"), record(2, None, 1, "BLOCKED REASON: Later note")
        )
        self.assertEqual(key_analysis([outcome]).blocked_reason, "Final reason")

    def test_earlier_retries_are_never_used(self):
        outcome = run_outcome(500, record(1, BLOCKED, 2, "BLOCKED REASON: Retry reason"), record(2, BLOCKED, 1))
        self.assertIsNone(key_analysis([outcome]).blocked_reason)

    def test_historical_blocked_reason_is_not_shown_for_a_newer_status(self):
        older = run_outcome(500, record(1, BLOCKED, 3, "BLOCKED REASON: Old outage"))
        for status in (PASSED, FAILED):
            with self.subTest(status=status):
                newer = run_outcome(501, record(2, status, 1, "BLOCKED REASON: stale label"))
                self.assertIsNone(key_analysis([older, newer]).blocked_reason)

    def test_latest_result_is_chosen_by_timestamp_not_run_id(self):
        newer_blocked = run_outcome(400, record(1, BLOCKED, 1, "BLOCKED REASON: Current outage"))
        older_passed = run_outcome(900, record(2, PASSED, 5))
        self.assertEqual(key_analysis([older_passed, newer_blocked]).blocked_reason, "Current outage")

    def test_unverified_blocked_final_has_no_reason(self):
        outcome = run_outcome(500, record(1, BLOCKED, 1, RAW), complete=False)
        self.assertIsNone(key_analysis([outcome]).blocked_reason)

    def test_disabled_extraction(self):
        outcome = run_outcome(500, record(1, BLOCKED, 1, RAW))
        self.assertIsNone(key_analysis([outcome], settings(enabled=False)).blocked_reason)

    def test_configured_headers_and_length(self):
        outcome = run_outcome(500, record(1, BLOCKED, 1, "Root cause: Vendor API rate limited for the suite"))
        self.assertIsNone(key_analysis([outcome]).blocked_reason)
        config = settings(section_headers=["Root cause"], max_length=12)
        self.assertEqual(key_analysis([outcome], config).blocked_reason, "Vendor API…")

    def test_reason_does_not_change_the_classification(self):
        records = [run_outcome(500 + i, record(i, BLOCKED, i + 1, RAW if i == 0 else None)) for i in range(3)]
        with_reason = key_analysis(records)
        without = key_analysis(records, settings(enabled=False))
        self.assertEqual(with_reason.blocked_reason, REASON)
        self.assertEqual(
            (with_reason.diagnosis, with_reason.confidence, with_reason.metrics),
            (without.diagnosis, without.confidence, without.metrics),
        )


class ReportOutputTests(unittest.TestCase):
    def sample(self):
        """100 blocked with a reason, 101 blocked without one, 102 passed after a blocked run."""
        return [
            run_outcome(500, record(1, BLOCKED, 1, RAW), case_id=100),
            run_outcome(510, record(2, BLOCKED, 1, "Blocked, see logs"), case_id=101),
            run_outcome(520, record(3, BLOCKED, 3, "BLOCKED REASON: Old outage"), case_id=102),
            run_outcome(521, record(4, PASSED, 1), case_id=102),
        ]

    def content(self, config=None):
        collected, analysis, config = analysis_of(self.sample(), config)
        return build_report(collected, analysis, config)

    @staticmethod
    def rows(sheet):
        headers = [name for name, _ in sheet.columns]
        return headers, {row[0].value: dict(zip(headers, row, strict=True)) for row in sheet.rows}

    def test_details_column_per_combination(self):
        headers, rows = self.rows(self.content().details)
        self.assertEqual([h for h in headers if h.endswith("Blocked Reason")],
                         [f"{app} | {platform}: Blocked Reason" for app in ("App A", "App B")
                          for platform in ("Platform X", "Platform Y")])  # fmt: skip
        self.assertEqual(headers[headers.index("App A | Platform X: Coverage") + 1], COLUMN)
        self.assertEqual((rows[100][COLUMN].value, rows[100][COLUMN].style), (REASON, None))
        self.assertEqual((rows[101][COLUMN].value, rows[101][COLUMN].style), ("Reason unavailable", "neutral"))
        self.assertIsNone(rows[102][COLUMN].value)  # latest status is Passed
        self.assertEqual(rows[100]["App A | Platform X: Latest Status"].value, "Blocked")
        self.assertIn(COLUMN, self.content().details.wrap)

    def test_overview_note_keeps_the_blocked_status(self):
        _, rows = self.rows(self.content().overview)
        cell = rows[100]["App A Platform X"]
        self.assertEqual(cell.value, "Blocked")
        self.assertEqual(cell.comment.splitlines()[-1], f"Blocked reason: {REASON}")
        self.assertEqual(rows[101]["App A Platform X"].comment.splitlines()[-1], "Blocked reason: Reason unavailable")
        self.assertNotIn("Blocked reason", rows[102]["App A Platform X"].comment)

    def test_disabled_extraction_omits_column_and_note(self):
        report = self.content(settings(enabled=False))
        self.assertFalse([name for name, _ in report.details.columns if name.endswith("Blocked Reason")])
        _, rows = self.rows(report.overview)
        self.assertNotIn("Blocked reason", rows[100]["App A Platform X"].comment)
        info = {row[0].value: row[1].value for row in report.info.rows}
        self.assertEqual(info["Blocked reasons (latest status Blocked, shown histories)"], "disabled")

    def test_report_info_counts_reasons(self):
        info = {row[0].value: row[1].value for row in self.content().info.rows}
        self.assertEqual(info["Blocked reasons (latest status Blocked, shown histories)"], "1 extracted, 1 unavailable")
        self.assertIn("Legend: Blocked Reason", info)

    def test_xml_has_the_column_and_no_raw_comment(self):
        collected, analysis, config = analysis_of(self.sample())
        xml = render_spreadsheet(collected, analysis, config)
        _, rows = table(xml)
        self.assertEqual(rows[0][COLUMN][0], REASON)
        for secret in ("SECRET-SESSION-42", "ci.example", "attachments", "Traceback", "RAW-COMMENT-TAIL",
                       "Automation run 77", "Old outage", "see logs"):  # fmt: skip
            self.assertNotIn(secret, xml)

    @unittest.skipUnless(openpyxl, "openpyxl is not installed")
    def test_xlsx_note_column_and_no_raw_comment(self):
        collected, analysis, config = analysis_of(self.sample())
        workbook = openpyxl.load_workbook(io.BytesIO(render_xlsx(collected, analysis, config)))
        overview, details = workbook["Stability Overview"], workbook["Stability Details"]
        headers = [cell.value for cell in overview[1]]
        cell = overview.cell(2, headers.index("App A Platform X") + 1)
        self.assertEqual(cell.value, "Blocked")
        self.assertIn(f"Blocked reason: {REASON}", cell.comment.text)
        headers = [cell.value for cell in details[1]]
        reason_cell = details.cell(2, headers.index(COLUMN) + 1)
        self.assertEqual(reason_cell.value, REASON)
        self.assertTrue(reason_cell.alignment.wrap_text)
        cells = [c for ws in workbook.worksheets for r in ws.iter_rows() for c in r]
        text = "\n".join(f"{c.value} {c.comment.text if c.comment else ''}" for c in cells)
        for secret in ("SECRET-SESSION-42", "ci.example", "attachments", "Traceback", "RAW-COMMENT-TAIL"):
            self.assertNotIn(secret, text)

    def test_hostile_reason_is_neutralized(self):
        outcome = run_outcome(500, record(1, BLOCKED, 1, "BLOCKED REASON:\n=cmd|' /C calc'!A0 broke the env"))
        collected, analysis, config = analysis_of([outcome])
        _, rows = table(render_spreadsheet(collected, analysis, config))
        self.assertTrue(rows[0][COLUMN][0].startswith("'="))


class SettingsTests(unittest.TestCase):
    def test_defaults(self):
        reason = fx.settings().report.blocked_reason
        self.assertEqual((reason.enabled, reason.section_headers, reason.max_length),
                         (True, ("BLOCKED REASON", "SKIPPED REASON"), 160))  # fmt: skip

    def test_values_are_normalized(self):
        reason = settings(section_headers=["  Root   cause ", "Root cause"], max_length=50).report.blocked_reason
        self.assertEqual((reason.section_headers, reason.max_length), (("Root cause",), 50))

    def test_invalid_values(self):
        from testrail_stability_analyzer.config import ConfigurationError

        for section, message in (
            ({"section_headers": []}, "report.blocked_reason.section_headers"),
            ({"section_headers": [""]}, "report.blocked_reason.section_headers"),
            ({"section_headers": ["A\nB"]}, "report.blocked_reason.section_headers"),
            ({"section_headers": "BLOCKED REASON"}, "report.blocked_reason.section_headers"),
            ({"max_length": 0}, "report.blocked_reason.max_length must be >= 1"),
            ({"max_length": 1001}, "report.blocked_reason.max_length must be <= 1000"),
            ({"enabled": "yes"}, "report.blocked_reason.enabled must be true or false"),
            ({"fallback": True}, "report.blocked_reason.fallback: unknown setting"),
        ):
            with self.subTest(section=section), self.assertRaises(ConfigurationError) as ctx:
                settings(**section)
            self.assertTrue(any(message in error for error in ctx.exception.errors), ctx.exception.errors)


if __name__ == "__main__":
    unittest.main()
