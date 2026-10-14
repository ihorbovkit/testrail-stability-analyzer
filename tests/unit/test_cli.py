import contextlib
import io
import tempfile
import unittest
from pathlib import Path
from unittest import mock

from testrail_stability_analyzer import cli

from . import fake_testrail as fx
from .fake_testrail import AUTOMATION_ISSUE, FAILED, NOW, PASSED, FakeTestRail, days_ago
from .fakes import FAKE_KEY, FAKE_USER
from .test_export_spreadsheet import info, table

try:
    import openpyxl
except ImportError:  # pragma: no cover - exercised only without the optional dependency
    openpyxl = None

CONFIG = """
[testrail]
base_url = "https://example.testrail.io"
project_id = 1

[http]
max_retries = 0

[[mappings.applications]]
name = "App A"
patterns = ['(?i)app-a']
sources = ["plan_name", "run_name"]

[[mappings.platforms]]
name = "Platform X"
patterns = ['(?i)plat-x']
sources = ["configuration", "run_name"]

[[mappings.platforms]]
name = "Platform Y"
patterns = ['(?i)plat-y']
sources = ["configuration", "run_name"]

[mappings.statuses]
automation_issue = ["custom_status1"]
not_applicable = ["custom_status2"]
"""

ENVIRON = {"TESTRAIL_USERNAME": FAKE_USER, "TESTRAIL_API_KEY": FAKE_KEY}
SENSITIVE_COMMENT = "Synthetic failure session_id=abc123 at https://assets.example.invalid/shot.png"
SENSITIVE_DEFECT = "SYNTH-4242"
PLAN_NAME = "App-A Confidential Regression"
RUN_NAME = "Confidential nightly"


def build_server():
    server = FakeTestRail()
    t = days_ago(3)
    server.add_plan(900, PLAN_NAME, created=days_ago(4))
    server.add_plan_run(
        900,
        fx.run(501, RUN_NAME, created=days_ago(4), config="Plat-X"),
        tests=[fx.test(1, 100, 501), fx.test(2, 101, 501)],
        results=[
            fx.result(11, 1, FAILED, t, comment=SENSITIVE_COMMENT, defects=SENSITIVE_DEFECT),
            fx.result(12, 1, AUTOMATION_ISSUE, t + 60, comment="Synthetic triage note"),
            fx.result(13, 2, PASSED, t),
        ],
    )
    server.add_plan_run(
        900,
        fx.run(502, RUN_NAME, created=days_ago(4), config="Unknown-Platform"),
        tests=[fx.test(3, 100, 502)],
        results=[fx.result(14, 3, PASSED, t)],
    )
    server.add_run(
        fx.run(601, "App-A smoke Plat-X", created=days_ago(1)),
        tests=[fx.test(5, 100, 601)],
        results=[fx.result(16, 5, PASSED, days_ago(1))],
    )
    return server


class CliTestCase(unittest.TestCase):
    def setUp(self):
        directory = tempfile.TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        self.config = Path(directory.name) / "config.toml"
        self.config.write_text(CONFIG)
        self.server = build_server()

    def run_cli(self, *argv, environ=ENVIRON, config=True, progress="off"):
        """Progress is off unless requested, so stderr holds only the command's own messages."""
        stdout, stderr = io.StringIO(), io.StringIO()
        args = list(argv) + (["--config", str(self.config)] if config else [])
        args += ["--progress", progress] if config and progress is not None else []
        code = cli.main(args, stdout=stdout, stderr=stderr, environ=environ, transport=self.server, clock=lambda: NOW)
        return code, stdout.getvalue(), stderr.getvalue()


class ArgumentValidationTests(CliTestCase):
    def test_missing_command_or_ids(self):
        for argv in ([], ["collect"], ["collect", "plan"], ["collect", "run"]):
            with self.subTest(argv=argv):
                code, _, stderr = self.run_cli(*argv, config=False)
                self.assertEqual(code, cli.EXIT_INVALID_CONFIGURATION)
                self.assertIn("usage:", stderr)
        self.assertEqual(self.server.calls, [])

    def test_invalid_numbers(self):
        for argv in (
            ["collect", "plan", "0"],
            ["collect", "plan", "abc"],
            ["collect", "run", "5", "--max-results", "-1"],
            ["collect", "run", "5", "--max-requests", "0"],
            ["collect", "plan", "900", "--run-id", "x"],
        ):
            with self.subTest(argv=argv):
                code, _, _ = self.run_cli(*argv, config=False)
                self.assertEqual(code, cli.EXIT_INVALID_CONFIGURATION)

    def test_invalid_configuration(self):
        code, _, stderr = self.run_cli("collect", "plan", "900", config=False, environ={})

        self.assertEqual(code, cli.EXIT_INVALID_CONFIGURATION)
        self.assertIn("testrail.base_url is required", stderr)
        self.assertIn("testrail.project_id is required", stderr)

    def test_inconsistent_overrides_are_validated(self):
        code, _, stderr = self.run_cli("collect", "plan", "900", "--window-days", "90", "--lookback-days", "30")

        self.assertEqual(code, cli.EXIT_INVALID_CONFIGURATION)
        self.assertIn("history.lookback_limit_days", stderr)

    def test_keychain_without_names_is_a_configuration_error(self):
        code, _, stderr = self.run_cli("collect", "plan", "900", "--credential-provider", "keychain")

        self.assertEqual(code, cli.EXIT_INVALID_CONFIGURATION)
        self.assertIn("credential_provider.keychain_service is required", stderr)
        self.assertEqual(self.server.calls, [])


class ExitCodeTests(CliTestCase):
    def test_complete_targeted_collection_succeeds(self):
        code, stdout, _ = self.run_cli("collect", "plan", "900", "--lookback-days", "0")

        self.assertEqual(code, cli.EXIT_OK, stdout)
        # Run 502 is unidentified (per-run key), so the selection does not limit it.
        self.assertIn("coverage: complete=1, selection_limited=2", stdout)
        self.assertIn("Collection: OK", stdout)

    def test_capped_collection_is_partial(self):
        code, stdout, _ = self.run_cli("collect", "plan", "900", "--max-results", "2")

        self.assertEqual(code, cli.EXIT_PARTIAL)
        self.assertIn("Result records fetched: 2 (cap: 2; cap reached: yes)", stdout)
        self.assertIn("Collection: PARTIAL", stdout)

    def test_some_invalid_ids_are_partial(self):
        self.server.errors["get_plan/999"] = (400, "Field :plan_id is not a valid test plan.")
        code, stdout, _ = self.run_cli("collect", "plan", "999", "900")

        self.assertEqual(code, cli.EXIT_PARTIAL)
        self.assertIn("Invalid selection: plan 999: plan does not exist or is not a valid ID (HTTP 400)", stdout)

    def test_only_invalid_ids_is_an_api_error(self):
        self.server.errors["get_run/999"] = (400, "Field :run_id is not a valid test run.")
        code, _, stderr = self.run_cli("collect", "run", "999")

        self.assertEqual(code, cli.EXIT_API_ERROR)
        self.assertIn("none of the selected plans or runs could be collected", stderr)

    def test_only_invalid_ids_write_no_report(self):
        self.server.errors["get_run/999"] = (400, "Field :run_id is not a valid test run.")
        report = self.config.parent / "report.xml"
        code, stdout, stderr = self.run_cli("collect", "run", "999", "--xml-report", str(report))

        self.assertEqual(code, cli.EXIT_API_ERROR)
        self.assertIn("none of the selected plans or runs could be collected", stderr)
        self.assertNotIn("report written", stdout)
        self.assertFalse(report.exists())
        metadata = ("get_sections/", "get_cases/", "get_suites/", "get_case_types")
        self.assertEqual([call for call in self.server.calls if call.startswith(metadata)], [])

    def test_authentication_failure_is_an_api_error(self):
        self.server.errors["get_statuses"] = (401, "Authentication failed: invalid or missing user/password.")
        code, stdout, stderr = self.run_cli("collect", "plan", "900")

        self.assertEqual(code, cli.EXIT_API_ERROR)
        self.assertIn("TestRailAuthenticationError (HTTP 401)", stderr)
        self.assertNotIn("Authentication failed", stderr)
        self.assertEqual(stdout, "")
        self.assertNotIn(FAKE_KEY, stderr)

    def test_server_error_text_is_never_printed(self):
        server_text = f"{SENSITIVE_COMMENT} defect {SENSITIVE_DEFECT} echoed key {FAKE_KEY} user {FAKE_USER}"
        self.server.errors["get_plan/900"] = (500, server_text)
        code, stdout, stderr = self.run_cli("collect", "plan", "900")

        self.assertEqual(code, cli.EXIT_API_ERROR)
        self.assertEqual(stderr, "TestRail API error: TestRailHTTPError (HTTP 500); response details are not shown\n")
        for sensitive in (SENSITIVE_COMMENT, "abc123", "assets.example.invalid", SENSITIVE_DEFECT, FAKE_KEY, FAKE_USER):
            self.assertNotIn(sensitive, stdout + stderr)

    def test_missing_credentials_are_reported_without_requests(self):
        code, _, stderr = self.run_cli("collect", "plan", "900", environ={})

        self.assertEqual(code, cli.EXIT_API_ERROR)
        self.assertIn("TESTRAIL_USERNAME", stderr)
        self.assertEqual(self.server.calls, [])

    def test_request_budget_exhaustion_is_partial(self):
        code, stdout, _ = self.run_cli("collect", "plan", "900", "--max-requests", "3")

        self.assertEqual(code, cli.EXIT_PARTIAL)
        self.assertIn("API requests: 3 (budget: 3)", stdout)
        self.assertIn("Request budget exhausted: yes", stdout)

    def test_sample_without_eligible_plan(self):
        code, stdout, _ = self.run_cli("collect", "sample", "--max-candidates", "1")
        self.assertEqual(code, cli.EXIT_OK)  # plan 900 is eligible

        self.server.plans.clear()
        code, stdout, _ = self.run_cli("collect", "sample")
        self.assertEqual(code, cli.EXIT_NO_TARGET)
        self.assertIn("No eligible Test Plan found", stdout)


class CollectionCommandTests(CliTestCase):
    def test_run_command_with_plan_option_is_a_mixed_selection(self):
        code, stdout, _ = self.run_cli("collect", "run", "601", "--plan-id", "900", "--lookback-days", "0")

        self.assertEqual(code, cli.EXIT_OK, stdout)
        self.assertIn("Selected plan IDs: 900", stdout)
        self.assertIn("Selected run IDs: 601", stdout)
        self.assertIn("Runs inspected: 3 (included 3)", stdout)

    def test_targeted_command_skips_project_wide_discovery(self):
        self.run_cli("collect", "plan", "900")

        self.assertNotIn("get_runs/1", self.server.calls)
        self.assertNotIn("get_plans/1", self.server.calls)
        self.assertEqual(self.server.calls[:2], ["get_statuses", "get_plan/900"])

    def test_sample_uses_validation_defaults(self):
        code, stdout, _ = self.run_cli("collect", "sample")

        self.assertIn("Sample selection: plan 900", stdout)
        self.assertIn("(cap: 15;", stdout)
        self.assertIn("(budget: 20)", stdout)
        self.assertEqual(self.server.calls.count("get_plan/900"), 1)  # fetched once, reused by the collector
        calls = zip(self.server.calls, self.server.params, strict=True)
        results_params = [params for call, params in calls if call.startswith("get_results")]
        self.assertEqual(results_params[0]["limit"], "15")

    def test_sample_defaults_can_be_overridden(self):
        _, stdout, _ = self.run_cli("collect", "sample", "--max-results", "0", "--max-requests", "50")

        self.assertIn("(cap: none;", stdout)
        self.assertIn("(budget: 50)", stdout)

    def test_selection_in_configuration_is_replaced_by_the_command_line(self):
        self.config.write_text(CONFIG + "\n[collection]\nrun_ids = [601]\n")
        _, stdout, _ = self.run_cli("collect", "plan", "900")

        self.assertIn("Selected run IDs: none", stdout)
        self.assertNotIn("get_run/601", self.server.calls)


class ProjectCommandTests(CliTestCase):
    def test_project_discovers_runs_project_wide(self):
        code, stdout, _ = self.run_cli("collect", "project")

        self.assertEqual(code, cli.EXIT_OK, stdout)
        self.assertIn("TestRail historical collection (project-wide, read-only)", stdout)
        self.assertNotIn("Selected plan IDs", stdout)
        self.assertIn("Plans inspected: 900", stdout)
        self.assertIn("Runs inspected: 3 (included 3)", stdout)
        self.assertIn("run 601 (standalone): included", stdout)
        # Default lookback: created in it, plus older open runs and plans.
        self.assertEqual(self.server.calls.count("get_runs/1"), 2)
        self.assertEqual(self.server.calls.count("get_plans/1"), 2)
        self.assertEqual(self.server.calls.count("get_plan/900"), 1)
        self.assertFalse(any(call.startswith("get_run/") for call in self.server.calls))
        self.assertTrue(all(call.startswith("get_") for call in self.server.calls))  # read-only endpoints only
        self.assertNotIn("targeted collection", stdout)

    def test_project_without_lookback_lists_once(self):
        code, stdout, _ = self.run_cli("collect", "project", "--lookback-days", "0")

        self.assertEqual(code, cli.EXIT_OK, stdout)
        self.assertEqual(self.server.calls[:3], ["get_statuses", "get_runs/1", "get_plans/1"])
        self.assertIn("coverage: complete=3", stdout)

    def test_project_follows_pagination(self):
        _, expected, _ = self.run_cli("collect", "project")
        self.server = build_server()
        self.server.page_size = 1
        code, stdout, _ = self.run_cli("collect", "project")

        self.assertEqual(code, cli.EXIT_OK, stdout)
        self.assertEqual(self.server.calls.count("get_tests/501"), 2)  # two tests, one per page
        self.assertEqual(self.server.calls.count("get_results_for_run/501"), 3)
        strip = lambda text: [line for line in text.splitlines() if not line.startswith("API requests")]  # noqa: E731
        self.assertEqual(strip(stdout), strip(expected))

    def test_selection_in_configuration_is_ignored(self):
        self.config.write_text(CONFIG + "\n[collection]\nplan_ids = [900]\nrun_ids = [601]\n")
        code, stdout, _ = self.run_cli("collect", "project")

        self.assertEqual(code, cli.EXIT_OK, stdout)
        self.assertIn("(project-wide, read-only)", stdout)
        self.assertNotIn("get_run/601", self.server.calls)
        self.assertIn("get_runs/1", self.server.calls)

    def test_suite_filter_applies(self):
        code, stdout, stderr = self.run_cli("collect", "project", "--suite-id", "99")

        self.assertEqual(code, cli.EXIT_NO_TARGET)
        self.assertIn("Runs inspected: 3 (included 0)", stdout)
        self.assertIn("run 601 (standalone): excluded_by_filter", stdout)
        self.assertIn("no eligible run found in the project; nothing was collected", stderr)

    def test_no_eligible_run_writes_no_report(self):
        self.server = FakeTestRail()
        report = self.config.parent / "report.xml"
        code, stdout, stderr = self.run_cli("collect", "project", "--xml-report", str(report))

        self.assertEqual(code, cli.EXIT_NO_TARGET)
        self.assertIn("Runs inspected: 0 (included 0)", stdout)
        self.assertIn("nothing was collected", stderr)
        self.assertNotIn("report written", stdout)
        self.assertFalse(report.exists())
        metadata = XmlReportCommandTests.METADATA
        self.assertEqual([call for call in self.server.calls if call.startswith(metadata)], [])

    def test_budget_exhausted_before_any_run_is_discovered(self):
        report = self.config.parent / "report.xml"
        code, stdout, stderr = self.run_cli("collect", "project", "--max-requests", "1", "--xml-report", str(report))

        self.assertEqual(code, cli.EXIT_PARTIAL)
        self.assertIn("Request budget exhausted: yes", stdout)
        self.assertIn("request budget exhausted during run discovery", stdout)
        self.assertIn("nothing was collected", stderr)
        self.assertFalse(report.exists())

    def test_budget_exhausted_during_discovery_keeps_discovered_runs(self):
        # get_statuses and the first get_runs listing succeed; the second listing exceeds the budget.
        report = self.config.parent / "report.xml"
        code, stdout, _ = self.run_cli("collect", "project", "--max-requests", "2", "--xml-report", str(report))

        self.assertEqual(code, cli.EXIT_PARTIAL)
        self.assertIn("Runs inspected: 1 (included 1)", stdout)
        self.assertIn("request budget exhausted during run discovery", stdout)
        self.assertIn("Collection: PARTIAL", stdout)
        self.assertTrue(report.exists())  # a partial report, disclosed as such
        values = {label: value for label, value, _ in info(report.read_text(encoding="utf-8"))}
        self.assertEqual(values["Collection status"], "PARTIAL")

    def test_budget_exhausted_during_collection_is_partial(self):
        code, stdout, _ = self.run_cli("collect", "project", "--max-requests", "8")

        self.assertEqual(code, cli.EXIT_PARTIAL)
        self.assertIn("Runs inspected: 3 (included 3)", stdout)
        self.assertIn("request budget exhausted during collection; coverage is partial", stdout)
        self.assertIn("Collection: PARTIAL", stdout)

    def test_listing_errors_are_sanitized(self):
        for endpoint, status, name in (
            ("get_plans/1", 500, "TestRailHTTPError (HTTP 500)"),
            ("get_runs/1", 403, "TestRailAuthorizationError (HTTP 403)"),
        ):
            with self.subTest(endpoint=endpoint):
                self.server = build_server()
                self.server.errors[endpoint] = (status, f"{SENSITIVE_COMMENT} echoed key {FAKE_KEY}")
                code, stdout, stderr = self.run_cli("collect", "project")

                self.assertEqual(code, cli.EXIT_API_ERROR)
                self.assertIn(name, stderr)
                for sensitive in (SENSITIVE_COMMENT, "abc123", FAKE_KEY, FAKE_USER):
                    self.assertNotIn(sensitive, stdout + stderr)

    def test_project_takes_no_ids(self):
        code, _, stderr = self.run_cli("collect", "project", "900")

        self.assertEqual(code, cli.EXIT_INVALID_CONFIGURATION)
        self.assertIn("unrecognized arguments", stderr)
        self.assertEqual(self.server.calls, [])


class MaxRunsOptionTests(CliTestCase):
    def overrides(self, *argv):
        return cli._overrides(cli.build_parser().parse_args(list(argv)))

    def test_overrides(self):
        self.assertEqual(self.overrides("collect", "project", "--max-runs", "20")["history"], {"max_runs": 20})
        self.assertNotIn("history", self.overrides("collect", "project"))
        for argv in (["collect", "plan", "900"], ["collect", "run", "5"], ["collect", "sample"], ["collect", "project"]):
            with self.subTest(argv=argv):
                self.assertIn("max_runs", self.overrides(*argv, "--max-runs", "3")["history"])

    def test_project_selection_overrides_are_empty(self):
        collection = self.overrides("collect", "project", "--window-days", "90")["collection"]
        self.assertEqual(collection, {"plan_ids": [], "run_ids": []})
        self.assertEqual(self.overrides("collect", "project", "--window-days", "90")["history"], {"window_days": 90})

    def test_targeted_overrides_are_unchanged(self):
        self.assertEqual(
            self.overrides("collect", "plan", "900", "--run-id", "5")["collection"], {"plan_ids": [900], "run_ids": [5]}
        )
        self.assertEqual(self.overrides("collect", "run", "5")["collection"], {"plan_ids": [], "run_ids": [5]})
        sample = self.overrides("collect", "sample")
        self.assertEqual(sample["collection"], {"max_results": cli.SAMPLE_MAX_RESULTS, "plan_ids": [], "run_ids": []})

    def test_invalid_values_are_rejected_before_requests(self):
        for value in ("0", "-1", "x"):
            with self.subTest(value=value):
                code, _, stderr = self.run_cli("collect", "project", "--max-runs", value)
                self.assertEqual(code, cli.EXIT_INVALID_CONFIGURATION)
                self.assertIn("usage:", stderr)
        self.assertEqual(self.server.calls, [])

    def test_validated_against_min_executions(self):
        self.config.write_text(CONFIG + "\n[analysis]\nmin_executions = 5\n")
        code, _, stderr = self.run_cli("collect", "project", "--max-runs", "3")

        self.assertEqual(code, cli.EXIT_INVALID_CONFIGURATION)
        self.assertIn("analysis.min_executions must be <= history.max_runs", stderr)
        self.assertEqual(self.server.calls, [])

    def test_max_runs_limits_the_analysis_window(self):
        # Case 100 ran in runs 501 and 601, both App A / Platform X.
        report = self.config.parent / "report.xml"
        executions = {}
        for max_runs in ("1", "20"):
            self.server = build_server()
            code, stdout, _ = self.run_cli(
                "collect", "project", "--lookback-days", "0", "--max-runs", max_runs, "--xml-report", str(report),
                "--overwrite-report",
            )  # fmt: skip
            self.assertEqual(code, cli.EXIT_OK, stdout)
            _, rows = table(report.read_text(encoding="utf-8"))
            row = next(row for row in rows if row["Test Case ID"][0] == "100")
            executions[max_runs] = row["App A | Platform X: Executions"][0]
        self.assertEqual(executions, {"1": "1", "20": "2"})


class SanitizedOutputTests(CliTestCase):
    def test_summary_contains_counts_but_no_sensitive_values(self):
        code, stdout, stderr = self.run_cli("collect", "plan", "900", "--lookback-days", "0")
        output = stdout + stderr

        for secret in (FAKE_KEY, FAKE_USER, SENSITIVE_COMMENT, "abc123", "assets.example.invalid", SENSITIVE_DEFECT,
                       "Synthetic triage note", PLAN_NAME, RUN_NAME, "Authorization"):
            self.assertNotIn(secret, output)
        self.assertIn("Project ID: 1", stdout)
        self.assertIn("Plans inspected: 900", stdout)
        self.assertIn("run 501 (plan 900): included; application=App A, platform=Platform X", stdout)
        self.assertIn("Unresolved classification: run 502 platform unresolved", stdout)
        self.assertIn("Test instances examined: 3", stdout)
        self.assertIn("Result records fetched: 4 (cap: none; cap reached: no)", stdout)
        self.assertIn("  with comments: 2", stdout)
        self.assertIn("  with linked defects: 1", stdout)
        self.assertIn("tests with retries or triage (more than one record in a run): 1", stdout)
        self.assertIn("Final status categories: automation_issue=1, passed=2", stdout)
        self.assertIn("groups: App A / Platform X=", stdout)
        self.assertIn("unidentified (grouped per run)=1", stdout)
        self.assertIn("API requests: 6", stdout)

    def test_summary_is_deterministic(self):
        first = self.run_cli("collect", "plan", "900")
        self.server = build_server()
        self.assertEqual(first, self.run_cli("collect", "plan", "900"))


class XmlReportCommandTests(CliTestCase):
    METADATA = ("get_sections/", "get_cases/", "get_suites/", "get_case_types")

    def setUp(self):
        super().setUp()
        self.server.cases[10] = [fx.payloads.case(100, section_id=2), fx.payloads.case(101, section_id=1)]
        self.server.sections[10] = [fx.payloads.section(1, "Area One"), fx.payloads.section(2, "Feature", parent_id=1)]
        self.report = self.config.parent / "report.xml"

    def metadata_calls(self):
        return [call for call in self.server.calls if call.startswith(self.METADATA)]

    def test_report_with_feature_groups(self):
        code, stdout, _ = self.run_cli("collect", "plan", "900", "--lookback-days", "0", "--xml-report", str(self.report))

        self.assertEqual(code, cli.EXIT_OK, stdout)
        self.assertIn(f"XML report written: {self.report}", stdout)
        self.assertIn("Feature / Test Groups: resolved=2; unknown: none", stdout)
        self.assertEqual(self.metadata_calls(), ["get_cases/1", "get_sections/1", "get_case_types"])
        self.assertIn("API requests: 9", stdout)
        xml = self.report.read_text(encoding="utf-8")
        self.assertIn("<Data ss:Type=\"String\">Feature / Test Group</Data>", xml)
        self.assertIn("<Data ss:Type=\"String\">Area One/Feature</Data>", xml)
        self.assertIn("<Data ss:Type=\"String\">Functional</Data>", xml)  # case type 3
        for sensitive in (SENSITIVE_COMMENT, "abc123", "assets.example.invalid", "Synthetic triage note",
                          PLAN_NAME, RUN_NAME, FAKE_KEY, FAKE_USER):
            self.assertNotIn(sensitive, xml)

    def test_collection_only_makes_no_metadata_requests(self):
        code, stdout, _ = self.run_cli("collect", "plan", "900", "--lookback-days", "0")

        self.assertEqual(code, cli.EXIT_OK)
        self.assertEqual(self.metadata_calls(), [])
        self.assertNotIn("Feature / Test Groups", stdout)
        self.assertNotIn("XML report written", stdout)

    def test_disabled_feature_groups(self):
        self.config.write_text(CONFIG + "\n[report.feature_group]\nenabled = false\n")
        code, _, _ = self.run_cli("collect", "plan", "900", "--xml-report", str(self.report))

        self.assertEqual(self.metadata_calls(), ["get_case_types"])
        names, _ = table(self.report.read_text(encoding="utf-8"))
        self.assertEqual(names[:4], ["Test Case ID", "Test Case Title", "Test Type", "Defects"])

    def test_budget_exhausted_by_metadata_keeps_the_collection_result(self):
        code, stdout, _ = self.run_cli(
            "collect", "plan", "900", "--lookback-days", "0", "--max-requests", "6", "--xml-report", str(self.report)
        )

        self.assertEqual(code, cli.EXIT_OK, stdout)  # collection itself completed
        self.assertIn("Request budget exhausted: no", stdout)
        self.assertIn("Feature / Test Groups: resolved=0; unknown: not_fetched=2", stdout)
        xml = self.report.read_text(encoding="utf-8")
        self.assertIn("Unknown (not fetched)", xml)
        self.assertIn("Test Type shows TestRail type IDs", xml)
        # Stability results are exported unchanged.
        self.assertIn('<Cell ss:StyleID="yellow"><Data ss:Type="String">Automation Issue</Data></Cell>', xml)

    def test_invalid_report_path_is_rejected_before_requests(self):
        missing = self.config.parent / "missing"
        for path in (missing / "report.xml", f"{missing}/"):
            with self.subTest(path=path):
                code, _, stderr = self.run_cli("collect", "plan", "900", "--xml-report", str(path))
                self.assertEqual(code, cli.EXIT_INVALID_CONFIGURATION)
                self.assertIn("invalid report path", stderr)
        self.assertEqual(self.server.calls, [])

    def test_same_report_file_via_another_path_is_rejected_before_requests(self):
        link = self.config.parent / "link.xml"
        link.symlink_to(self.report)
        hard_link = self.config.parent / "hard.xml"
        self.report.write_text("existing")
        hard_link.hardlink_to(self.report)
        for other in (self.report.absolute(), self.config.parent / "." / "report.xml", link, hard_link):
            with self.subTest(other=other):
                code, _, stderr = self.run_cli(
                    "collect", "plan", "900", "--xml-report", str(self.report), "--xlsx-report", str(other)
                )
                self.assertEqual(code, cli.EXIT_INVALID_CONFIGURATION)
                self.assertIn("--xml-report and --xlsx-report must be different files", stderr)
        self.assertEqual(self.server.calls, [])


@unittest.skipUnless(openpyxl, "openpyxl is not installed")
class XlsxReportCommandTests(CliTestCase):
    def setUp(self):
        super().setUp()
        self.server.cases[10] = [fx.payloads.case(100, section_id=2), fx.payloads.case(101, section_id=1)]
        self.server.sections[10] = [fx.payloads.section(1, "Area One"), fx.payloads.section(2, "Feature", parent_id=1)]
        self.xlsx = self.config.parent / "report.xlsx"
        self.xml = self.config.parent / "report.xml"

    def metadata_calls(self):
        return [call for call in self.server.calls if call.startswith(XmlReportCommandTests.METADATA)]

    def test_xlsx_report(self):
        code, stdout, _ = self.run_cli(
            "collect", "plan", "900", "--lookback-days", "0", "--xlsx-report", str(self.xlsx)
        )

        self.assertEqual(code, cli.EXIT_OK, stdout)
        self.assertIn(f"XLSX report written: {self.xlsx}", stdout)
        workbook = openpyxl.load_workbook(self.xlsx)
        self.assertEqual(workbook.sheetnames, ["Stability Overview", "Stability Details", "Report Info"])
        overview = workbook["Stability Overview"]
        self.assertEqual(
            [cell.value for cell in overview[1]],
            ["Test Case ID", "Test Case Title", "Feature / Test Group", "Test Type", "Defects",
             "App A Platform X", "App A Platform Y"],
        )  # fmt: skip
        self.assertEqual(overview["B2"].value, "Case 100")  # the title from get_tests
        self.assertEqual(overview["C2"].value, "Area One/Feature")
        self.assertEqual(overview["F2"].value, "Automation Issue")
        texts = " ".join(str(c.value) for ws in workbook.worksheets for row in ws.iter_rows() for c in row)
        for sensitive in (SENSITIVE_COMMENT, "abc123", "Synthetic triage note", PLAN_NAME, RUN_NAME, FAKE_KEY):
            self.assertNotIn(sensitive, texts)

    def test_xml_and_xlsx_share_one_metadata_fetch(self):
        code, stdout, _ = self.run_cli(
            "collect", "plan", "900", "--lookback-days", "0",
            "--xml-report", str(self.xml), "--xlsx-report", str(self.xlsx),
        )  # fmt: skip

        self.assertEqual(code, cli.EXIT_OK, stdout)
        self.assertEqual(self.metadata_calls(), ["get_cases/1", "get_sections/1", "get_case_types"])
        self.assertIn("API requests: 9", stdout)
        self.assertIn("XML report written", stdout)
        self.assertIn("XLSX report written", stdout)
        xml_names, xml_rows = table(self.xml.read_text(encoding="utf-8"))
        details = openpyxl.load_workbook(self.xlsx)["Stability Details"]
        self.assertEqual([cell.value for cell in details[1]], xml_names)
        self.assertEqual(details.max_row - 1, len(xml_rows))

    def test_invalid_xlsx_paths(self):
        for argv in (
            ["--xlsx-report", str(self.config.parent / "missing" / "r.xlsx")],
            ["--xml-report", str(self.xml), "--xlsx-report", str(self.xml)],
        ):
            with self.subTest(argv=argv):
                code, _, stderr = self.run_cli("collect", "plan", "900", *argv)
                self.assertEqual(code, cli.EXIT_INVALID_CONFIGURATION)
                self.assertIn("invalid report path", stderr)
        self.assertEqual(self.server.calls, [])

    def test_project_xlsx_report(self):
        code, stdout, _ = self.run_cli("collect", "project", "--window-days", "90", "--max-runs", "20",
                                       "--xlsx-report", str(self.xlsx))  # fmt: skip

        self.assertEqual(code, cli.EXIT_OK, stdout)
        self.assertIn(f"XLSX report written: {self.xlsx}", stdout)
        workbook = openpyxl.load_workbook(self.xlsx)
        self.assertEqual(workbook.sheetnames, ["Stability Overview", "Stability Details", "Report Info"])
        info = {row[0].value: row[1].value for row in workbook["Report Info"].iter_rows() if len(row) > 1}
        self.assertEqual(info["Collection mode"], "project-wide")
        self.assertEqual(self.metadata_calls(), ["get_cases/1", "get_sections/1", "get_case_types"])
        texts = " ".join(str(c.value) for ws in workbook.worksheets for row in ws.iter_rows() for c in row)
        for sensitive in (SENSITIVE_COMMENT, "abc123", "Synthetic triage note", PLAN_NAME, RUN_NAME, FAKE_KEY):
            self.assertNotIn(sensitive, texts)

    def test_collection_only_still_makes_no_metadata_requests(self):
        self.run_cli("collect", "plan", "900")
        self.assertEqual(self.metadata_calls(), [])


class MissingOpenpyxlTests(CliTestCase):
    def test_missing_openpyxl_is_reported_before_requests(self):
        with mock.patch.object(cli, "xlsx_available", return_value=False):
            path = self.config.parent / "report.xlsx"
            code, _, stderr = self.run_cli("collect", "plan", "900", "--xlsx-report", str(path))

        self.assertEqual(code, cli.EXIT_INVALID_CONFIGURATION)
        self.assertIn('pip install "openpyxl>=3.1,<4"', stderr)
        self.assertEqual(self.server.calls, [])
        self.assertFalse((self.config.parent / "report.xlsx").exists())


class ProgressOutputTests(CliTestCase):
    def progress_lines(self, stderr):
        return [line.split("] ", 1)[1] for line in stderr.splitlines() if line.startswith("[")]

    def test_collection_without_reports_shows_three_phases(self):
        code, stdout, stderr = self.run_cli("collect", "plan", "900", progress="plain")

        self.assertEqual(code, cli.EXIT_OK)
        phases = [line for line in self.progress_lines(stderr) if ": " in line and line.startswith("Phase")]
        self.assertEqual(
            phases,
            [
                "Phase 1/3: Configuration and authentication (0 completed, 2 remaining after this)",
                "Phase 2/3: Run discovery (selected plans and runs) (1 completed, 1 remaining after this)",
                "Phase 3/3: Execution result collection (2 completed, 0 remaining after this)",
            ],
        )
        self.assertRegex(stderr.splitlines()[-1], r"Completed all 3 phases in \d+s; API requests \d+/1000$")
        self.assertNotIn("\r", stderr)

    def test_reports_add_metadata_analysis_and_generation_phases(self):
        report = self.config.parent / "report.xml"
        code, _, stderr = self.run_cli("collect", "project", "--xml-report", str(report), progress="plain")

        self.assertEqual(code, cli.EXIT_OK)
        lines = self.progress_lines(stderr)
        names = [line.split(": ", 1)[1] for line in lines if line.startswith("Phase ") and ": " in line]
        self.assertEqual(
            [name.rsplit(" (", 1)[0] for name in names],
            [
                "Configuration and authentication",
                "Run discovery (project-wide)",
                "Execution result collection",
                "Report metadata (feature groups, case types)",
                "Stability analysis",
                "Report generation (XML)",
            ],
        )
        self.assertIn("Completed all 6 phases", stderr)

    def test_stdout_and_exit_codes_are_unchanged_by_progress(self):
        self.server.errors["get_plan/999"] = (400, "Field :plan_id is not a valid test plan.")
        scenarios = [
            ("collect", "plan", "900"),
            ("collect", "plan", "999", "900"),
            ("collect", "plan", "900", "--max-requests", "3"),
            ("collect", "project", "--suite-id", "99"),
            ("collect", "sample"),
        ]
        for argv in scenarios:
            with self.subTest(argv=argv):
                quiet = self.run_cli(*argv)
                self.server = build_server()
                self.server.errors["get_plan/999"] = (400, "Field :plan_id is not a valid test plan.")
                plain = self.run_cli(*argv, progress="plain")
                self.assertEqual(plain[:2], quiet[:2])
                # The command's own stderr messages follow the progress output unchanged.
                self.assertTrue(plain[2].endswith(quiet[2]))

    def test_partial_collection_is_not_reported_as_completed(self):
        code, _, stderr = self.run_cli("collect", "plan", "900", "--max-requests", "3", progress="plain")

        self.assertEqual(code, cli.EXIT_PARTIAL)
        self.assertIn("Request budget exhausted (3/3)", stderr)
        self.assertIn("Phase 3/3 ended after 0s (request budget exhausted)", stderr)
        final = stderr.splitlines()[-1]
        self.assertIn("Finished with PARTIAL coverage", final)
        self.assertIn("request budget exhausted", final)
        self.assertNotIn("Completed", stderr)

    def test_failed_collection_names_the_phase(self):
        self.server.errors["get_plan/900"] = (500, SENSITIVE_COMMENT)
        code, stdout, stderr = self.run_cli("collect", "plan", "900", progress="plain")

        self.assertEqual(code, cli.EXIT_API_ERROR)
        lines = stderr.splitlines()
        self.assertIn("Failed during phase 2/3 (Run discovery (selected plans and runs))", lines[-2])
        self.assertIn("TestRail API error TestRailHTTPError (HTTP 500)", lines[-2])
        self.assertEqual(lines[-1], "TestRail API error: TestRailHTTPError (HTTP 500); response details are not shown")
        self.assertNotIn("Completed", stderr)
        self.assertNotIn("abc123", stderr)

    def test_invalid_configuration_fails_in_the_first_phase(self):
        code, _, stderr = self.run_cli(
            "collect", "plan", "900", "--window-days", "90", "--lookback-days", "30", progress="plain"
        )

        self.assertEqual(code, cli.EXIT_INVALID_CONFIGURATION)
        self.assertIn("Failed during phase 1/3 (Configuration and authentication)", stderr)
        self.assertIn("invalid configuration", stderr)

    def test_nothing_to_collect_stops_without_completion(self):
        self.server.plans.clear()
        code, _, stderr = self.run_cli("collect", "sample", progress="plain")

        self.assertEqual(code, cli.EXIT_NO_TARGET)
        self.assertIn("Stopped after phase 2/3 (Run discovery (sample plan selection))", stderr)
        self.assertNotIn("Completed", stderr)

    def test_progress_never_shows_names_comments_or_credentials(self):
        report = self.config.parent / "report.xml"
        _, _, stderr = self.run_cli("collect", "project", "--xml-report", str(report), progress="tty")

        sensitive_values = (PLAN_NAME, RUN_NAME, SENSITIVE_COMMENT, SENSITIVE_DEFECT, FAKE_KEY, FAKE_USER)
        for sensitive in (*sensitive_values, "Authorization"):
            self.assertNotIn(sensitive, stderr)

    def test_auto_mode_follows_the_terminal(self):
        class Terminal(io.StringIO):
            def isatty(self):
                return True

        def reporter(stream, *options, environ=None):
            args = cli.build_parser().parse_args(["collect", "project", *options])
            return cli._progress(args, stream, environ or {})

        self.assertFalse(reporter(io.StringIO()).interactive)  # redirected output or CI
        self.assertTrue(reporter(Terminal()).interactive)
        self.assertFalse(reporter(Terminal(), environ={"TERM": "dumb"}).interactive)
        self.assertFalse(reporter(Terminal(), "--progress", "plain").interactive)
        self.assertTrue(reporter(io.StringIO(), "--progress", "tty").interactive)
        self.assertIsInstance(reporter(Terminal(), "--progress", "off"), cli.NullProgress)

    def test_progress_off_prints_no_progress(self):
        _, _, stderr = self.run_cli("collect", "plan", "900", "--max-requests", "3")
        self.assertEqual(stderr, "")

    def test_invalid_progress_options(self):
        for options in (["--progress", "fancy"], ["--progress-interval", "0"]):
            with self.subTest(options=options):
                code, _, stderr = self.run_cli("collect", "plan", "900", *options)
                self.assertEqual(code, cli.EXIT_INVALID_CONFIGURATION)
                self.assertIn("usage:", stderr)


class ReportFileNameTests(CliTestCase):
    STAMP = "20260601-000000"  # NOW in UTC

    def setUp(self):
        super().setUp()
        self.out = self.config.parent / "reports"
        self.out.mkdir()

    def written(self, stdout):
        return [line.split(": ", 1)[1] for line in stdout.splitlines() if " report written: " in line]

    def test_automatic_names_per_collection_mode(self):
        cases = [
            (("collect", "project"), "stability-project-p1-60d"),
            (("collect", "project", "--suite-id", "10", "--window-days", "90"), "stability-project-p1-s10-90d"),
            (("collect", "plan", "900"), "stability-plan-900-p1-60d"),
            (("collect", "plan", "900", "--run-id", "601"), "stability-plan-900-run-601-p1-60d"),
            (("collect", "run", "601"), "stability-run-601-p1-60d"),
            (("collect", "sample"), "stability-sample-plan-900-p1-60d"),
        ]
        for argv, prefix in cases:
            with self.subTest(argv=argv):
                self.server = build_server()
                code, stdout, _ = self.run_cli(*argv, "--xml-report", f"{self.out}/")
                self.assertIn(code, (cli.EXIT_OK, cli.EXIT_PARTIAL), stdout)
                (path,) = self.written(stdout)
                self.assertEqual(path, str(self.out / f"{prefix}-{self.STAMP}.xml"))
                self.assertTrue(Path(path).is_file())

    def test_existing_directory_without_trailing_separator(self):
        code, stdout, _ = self.run_cli("collect", "plan", "900", "--xml-report", str(self.out))

        self.assertEqual(code, cli.EXIT_OK)
        self.assertEqual(self.written(stdout), [str(self.out / f"stability-plan-900-p1-60d-{self.STAMP}.xml")])

    def test_reports_in_the_same_second_get_unique_names(self):
        first = self.run_cli("collect", "plan", "900", "--xml-report", f"{self.out}/")
        self.server = build_server()
        second = self.run_cli("collect", "plan", "900", "--xml-report", f"{self.out}/", "--overwrite-report")

        base = self.out / f"stability-plan-900-p1-60d-{self.STAMP}"
        self.assertEqual(self.written(first[1]), [f"{base}.xml"])
        self.assertEqual(self.written(second[1]), [f"{base}-2.xml"])
        self.assertEqual(sorted(p.name for p in self.out.iterdir()), [f"{base.name}-2.xml", f"{base.name}.xml"])

    @unittest.skipUnless(openpyxl, "openpyxl is not installed")
    def test_both_formats_share_the_basename(self):
        code, stdout, _ = self.run_cli(
            "collect", "project", "--xlsx-report", f"{self.out}/", "--xml-report", f"{self.out}/"
        )

        self.assertEqual(code, cli.EXIT_OK, stdout)
        base = str(self.out / f"stability-project-p1-60d-{self.STAMP}")
        self.assertEqual(self.written(stdout), [f"{base}.xml", f"{base}.xlsx"])
        openpyxl.load_workbook(f"{base}.xlsx")

    def test_explicit_file_path_is_kept(self):
        report = self.out / "my-report.xml"
        code, stdout, _ = self.run_cli("collect", "plan", "900", "--xml-report", str(report))

        self.assertEqual(code, cli.EXIT_OK)
        self.assertEqual(self.written(stdout), [str(report)])
        self.assertEqual([p.name for p in self.out.iterdir()], ["my-report.xml"])

    def test_relative_paths_are_printed_absolute(self):
        with contextlib.chdir(self.config.parent):
            code, stdout, _ = self.run_cli("collect", "plan", "900", "--xml-report", "reports/")
        self.assertEqual(code, cli.EXIT_OK)
        (path,) = self.written(stdout)
        self.assertTrue(Path(path).is_absolute())
        self.assertTrue(Path(path).is_file())

    def test_existing_file_is_rejected_before_requests(self):
        report = self.out / "my-report.xml"
        report.write_text("previous report")
        code, stdout, stderr = self.run_cli("collect", "plan", "900", "--xml-report", str(report))

        self.assertEqual(code, cli.EXIT_INVALID_CONFIGURATION)
        self.assertIn(f"invalid report path: {report} already exists; pass --overwrite-report to replace it", stderr)
        self.assertEqual(self.server.calls, [])
        self.assertEqual((stdout, report.read_text()), ("", "previous report"))

    def test_overwrite_report_replaces_an_existing_file(self):
        report = self.out / "my-report.xml"
        report.write_text("previous report")
        code, stdout, _ = self.run_cli("collect", "plan", "900", "--xml-report", str(report), "--overwrite-report")

        self.assertEqual(code, cli.EXIT_OK)
        self.assertEqual(self.written(stdout), [str(report)])
        self.assertIn("<Workbook", report.read_text(encoding="utf-8"))
        self.assertEqual([p.name for p in self.out.iterdir()], ["my-report.xml"])

    def test_file_created_during_collection_is_not_overwritten(self):
        report = self.out / "my-report.xml"
        original = self.server.get

        def get(url, headers, **kwargs):
            if not report.exists():
                report.write_text("written by someone else")
            return original(url, headers, **kwargs)

        self.server.get = get
        code, stdout, stderr = self.run_cli("collect", "plan", "900", "--xml-report", str(report))

        self.assertEqual(code, cli.EXIT_UNEXPECTED)
        self.assertIn("could not write the XML report", stderr)
        self.assertIn("already exists", stderr)
        self.assertEqual(report.read_text(), "written by someone else")
        self.assertIn("Collection: OK", stdout)  # the summary is still printed

    def test_invalid_output_directories_make_no_requests(self):
        missing = self.config.parent / "missing"
        file = self.out / "file.xml"
        file.write_text("x")
        for raw in (f"{missing}/", str(missing / "r.xml"), f"{file}/"):
            with self.subTest(raw=raw):
                code, _, stderr = self.run_cli("collect", "project", "--xml-report", raw)
                self.assertEqual(code, cli.EXIT_INVALID_CONFIGURATION)
                self.assertIn("does not exist", stderr)
        self.assertEqual(self.server.calls, [])

    def test_help_documents_report_options(self):
        stdout = io.StringIO()
        with self.assertRaises(SystemExit), contextlib.redirect_stdout(stdout):
            cli.build_parser().parse_args(["collect", "project", "--help"])
        help_text = " ".join(stdout.getvalue().split())
        self.assertIn("--overwrite-report", help_text)
        self.assertIn("unique timestamped name", help_text)


if __name__ == "__main__":
    unittest.main()
