"""Native XLSX workbook, validated by reading it back with openpyxl. Skipped when openpyxl is not installed."""

import dataclasses
import io
import sys
import unittest
import zipfile
from datetime import datetime
from types import MappingProxyType
from unittest import mock

from testrail_stability_analyzer.analysis import analyze
from testrail_stability_analyzer.collection import CaseGroup, FeatureGroupState
from testrail_stability_analyzer.export import XlsxUnavailableError, render_xlsx, xlsx_available

from . import analysis_fixtures as af
from . import fake_testrail as fx
from .analysis_fixtures import outcome, rec
from .fake_testrail import FAILED, PASSED
from .report_fixtures import CASE_TYPES, HOSTILE_DEFECT, sample_inputs

try:
    import openpyxl
except ImportError:  # pragma: no cover - exercised only without the optional dependency
    openpyxl = None


def render(**kwargs):
    collection, analysis, settings = sample_inputs(**kwargs)
    return render_xlsx(collection, analysis, settings, case_types=CASE_TYPES, notes=("synthetic report note",))


def by_header(ws):
    headers = [cell.value for cell in ws[1]]
    return headers, {row[0].value: dict(zip(headers, row, strict=True)) for row in ws.iter_rows(min_row=2)}


def fill(cell):
    return cell.fill.fgColor.rgb[-6:] if cell.fill.fill_type == "solid" else None


@unittest.skipUnless(openpyxl, "openpyxl is not installed")
class WorkbookTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.data = render()
        cls.workbook = openpyxl.load_workbook(io.BytesIO(cls.data))

    def test_valid_package_without_formulas_or_external_links(self):
        with zipfile.ZipFile(io.BytesIO(self.data)) as package:
            self.assertIsNone(package.testzip())
            names = package.namelist()
        self.assertIn("xl/workbook.xml", names)
        self.assertFalse([name for name in names if "externalLink" in name])
        for ws in self.workbook.worksheets:
            self.assertEqual(ws._hyperlinks, [])
            for row in ws.iter_rows():
                for cell in row:
                    self.assertNotEqual(cell.data_type, "f", cell.coordinate)

    def test_worksheet_order(self):
        self.assertEqual(self.workbook.sheetnames, ["Stability Overview", "Stability Details", "Report Info"])

    def test_overview_columns_rows_and_statuses(self):
        headers, rows = by_header(self.workbook["Stability Overview"])
        self.assertEqual(headers[6:], ["App A Platform X", "App A Platform Y", "App B Platform X", "App B Platform Y"])
        self.assertEqual(list(rows), list(range(100, 109)))  # one row per case, numeric IDs
        cases = {
            100: ("Passed", "C6EFCE"),
            101: ("Failed", "FFC7CE"),
            102: ("Passed", "FFEB9C"),
            104: ("Not Applicable", "EDEDED"),
            106: ("Partial", "EDEDED"),
            108: ("No History", "EDEDED"),
        }
        for case_id, (label, colour) in cases.items():
            with self.subTest(case_id=case_id):
                cell = rows[case_id]["App A Platform X"]
                self.assertEqual((cell.value, fill(cell)), (label, colour))
        note = rows[100]["App A Platform Y"].comment
        self.assertTrue(note.text.startswith("Diagnosis: Insufficient History\n"))

    def test_case_columns(self):
        _, rows = by_header(self.workbook["Stability Overview"])
        self.assertEqual(rows[100]["Feature / Test Group"].value, "Area Two/Feature_Alpha")
        self.assertEqual(rows[101]["Feature / Test Group"].value, "R&D <Area>")
        self.assertEqual(rows[102]["Feature / Test Group"].value, "Unknown (case not found)")
        self.assertIsNone(rows[101]["Class Name"].value)  # the field is not set on its tests
        self.assertEqual(rows[100]["Test Type"].value, "Functional")
        defects = rows[101]["Defects"].value
        # Identifiers are kept as entered; only a leading formula character of the whole cell is neutralized.
        self.assertEqual(sorted(defects.split(", ")), [HOSTILE_DEFECT, "BUG-1", "BUG-2"])
        self.assertFalse(defects.startswith("="))
        self.assertTrue(rows[101]["Defects"].alignment.wrap_text)
        self.assertIsNone(rows[105]["Defects"].value)

    def test_class_name_from_the_newest_test(self):
        tests = MappingProxyType({"custom_class": "pkg.module.TestClass"})
        base = outcome(500, 100, rec(1, PASSED, 1))
        collection = af.report([dataclasses.replace(base, test=dataclasses.replace(base.test, custom_fields=tests))])
        settings = fx.settings(mappings={"custom_fields": {"class_name": "custom_class"}})
        data = render_xlsx(collection, analyze(collection, settings), settings)
        _, rows = by_header(openpyxl.load_workbook(io.BytesIO(data))["Stability Overview"])
        self.assertEqual(rows[100]["Class Name"].value, "pkg.module.TestClass")
        self.assertTrue(rows[100]["Class Name"].alignment.wrap_text)

    def test_formula_injection_is_neutralized(self):
        outcomes = [outcome(500, 100, rec(1, FAILED, 1, defects="=1+1"))]
        groups = {100: CaseGroup(100, FeatureGroupState.RESOLVED, 10, 1, ("@SUM(A1)", "-x"))}
        collection = af.report(outcomes, case_groups=groups)
        settings = fx.settings()
        data = render_xlsx(collection, analyze(collection, settings), settings)
        _, rows = by_header(openpyxl.load_workbook(io.BytesIO(data))["Stability Overview"])
        self.assertEqual(rows[100]["Defects"].value, "'=1+1")
        self.assertEqual(rows[100]["Defects"].data_type, "s")
        self.assertEqual(rows[100]["Feature / Test Group"].value, "'@SUM(A1)/-x")

    def test_details_numeric_percentages_and_dates(self):
        headers, rows = by_header(self.workbook["Stability Details"])
        self.assertEqual(len(headers), 6 + 4 * 10)
        row = rows[102]
        rate = row["App A | Platform X: Pass Rate"]
        self.assertEqual((rate.value, rate.number_format), (0.6, "0.0%"))
        latest = row["App A | Platform X: Latest Result (UTC)"]
        self.assertIsInstance(latest.value, datetime)
        self.assertEqual(latest.number_format, "yyyy-mm-dd hh:mm")
        self.assertEqual(row["App A | Platform X: Executions"].value, 5)
        self.assertEqual(row["App A | Platform X: Diagnosis"].value, "Inconsistent")
        self.assertEqual(fill(row["App A | Platform X: Diagnosis"]), "FFEB9C")
        self.assertEqual(rows[100]["App A | Platform Y: Diagnosis"].value, "Insufficient History")
        self.assertEqual(rows[106]["App A | Platform X: Coverage"].value, "Partial")
        self.assertEqual(rows[106]["App A | Platform X: Diagnosis"].value, "Partial")

    def test_frozen_headers_filters_and_widths(self):
        for name, frozen in (("Stability Overview", "B2"), ("Stability Details", "B2"), ("Report Info", "A2")):
            ws = self.workbook[name]
            with self.subTest(sheet=name):
                self.assertEqual(ws.freeze_panes, frozen)
                self.assertEqual(ws.auto_filter.ref, f"A1:{ws.cell(1, ws.max_column).column_letter}{ws.max_row}")
                self.assertTrue(all(cell.font.bold for cell in ws[1]))
                self.assertGreater(ws.column_dimensions["A"].width, 10)
        overview = self.workbook["Stability Overview"]
        self.assertEqual(overview.max_column, 10)  # compact: no per-metric columns

    def test_report_info(self):
        ws = self.workbook["Report Info"]
        items = [(row[0].value, row[1].value) for row in ws.iter_rows(min_row=2)]
        labels = [label for label, _ in items]
        self.assertIn("Overview statuses (shown histories)", labels)
        self.assertIn(("Limitation", "synthetic report note"), items)
        self.assertIn("Legend: Confidence", labels)
        self.assertIsInstance(dict(items)["Collected at (UTC)"], datetime)

    def test_sensitive_content_is_not_exported(self):
        text = "\n".join(
            str(cell.value) + str(cell.comment.text if cell.comment else "")
            for ws in self.workbook.worksheets
            for row in ws.iter_rows()
            for cell in row
        )
        for secret in ("Authorization", "session_id", "custom_steps"):
            self.assertNotIn(secret, text)


class MissingDependencyTests(unittest.TestCase):
    def test_clear_error_without_openpyxl(self):
        collection, analysis, settings = sample_inputs()
        with mock.patch.dict(sys.modules, {"openpyxl": None}):
            self.assertFalse(xlsx_available())
            with self.assertRaises(XlsxUnavailableError) as ctx:
                render_xlsx(collection, analysis, settings)
        self.assertIn('pip install "openpyxl>=3.1,<4"', str(ctx.exception))


if __name__ == "__main__":
    unittest.main()
