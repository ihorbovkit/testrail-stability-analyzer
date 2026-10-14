"""Stability Overview background = historical Pass Rate band; text = latest status. Synthetic data only."""

import dataclasses
import io
import re
import unittest

from testrail_stability_analyzer.analysis import analyze
from testrail_stability_analyzer.analysis.models import Diagnosis, Rate
from testrail_stability_analyzer.config.settings import StatusCategory
from testrail_stability_analyzer.export import build_report, render_spreadsheet
from testrail_stability_analyzer.export import overview_cell
from testrail_stability_analyzer.export.report_data import (
    NO_HISTORY_LABEL,
    OVERVIEW_GREEN,
    OVERVIEW_NEUTRAL,
    OVERVIEW_RED,
    OVERVIEW_YELLOW,
    overview_style,
)

from . import analysis_fixtures as af
from . import fake_testrail as fx
from .fake_testrail import AUTOMATION_ISSUE, BLOCKED, FAILED, NOT_APPLICABLE, PASSED

try:
    import openpyxl

    from testrail_stability_analyzer.export import render_xlsx
except ImportError:  # pragma: no cover - exercised only without the optional dependency
    openpyxl = None

COLUMN = "App A Platform X"
MAX_RUNS = 30


def settings(min_executions=5):
    return fx.settings(history={"max_runs": MAX_RUNS}, analysis={"min_executions": min_executions})


def inputs(statuses, *, min_executions=5, complete=True):
    """(collection, analysis, settings) for one App A / Platform X history; statuses are newest first."""
    outcomes = af.runs(statuses)
    if not complete:
        outcomes[-1] = dataclasses.replace(outcomes[-1], results_complete=False)
    collection = af.report(outcomes, hist=af.histories(outcomes, max_runs=MAX_RUNS))
    config = settings(min_executions)
    return collection, analyze(collection, config), config


def overview(statuses, **kwargs):
    collection, analysis, config = inputs(statuses, **kwargs)
    sheet = build_report(collection, analysis, config).overview
    headers = [name for name, _ in sheet.columns]
    (row,) = sheet.rows
    return dict(zip(headers, row, strict=True))[COLUMN], analysis.analyses[0]


class ThresholdTests(unittest.TestCase):
    """Exact fractions at and just below each threshold (rates far beyond a realistic run count)."""

    def style(self, passed, conclusive, *, min_executions=5):
        _, analysis, _ = inputs([PASSED] * 5)
        key = analysis.analyses[0]
        metrics = dataclasses.replace(key.metrics, pass_rate=Rate(passed, conclusive))
        return overview_style(dataclasses.replace(key, metrics=metrics), min_executions)

    def test_boundaries(self):
        cases = [
            (9, 10, OVERVIEW_GREEN),  # exactly 90%
            (900_000, 1_000_000, OVERVIEW_GREEN),
            (899_999, 1_000_000, OVERVIEW_YELLOW),  # 89.9999% displays as 90.0% but is below
            (5, 10, OVERVIEW_YELLOW),  # exactly 50%
            (499_999, 1_000_000, OVERVIEW_RED),  # 49.9999% displays as 50.0% but is below
            (0, 5, OVERVIEW_RED),
            (5, 5, OVERVIEW_GREEN),
        ]
        for passed, conclusive, expected in cases:
            with self.subTest(rate=f"{passed}/{conclusive}"):
                self.assertEqual(self.style(passed, conclusive), expected)

    def test_without_conclusive_executions_or_with_too_few(self):
        self.assertEqual(self.style(0, 0), OVERVIEW_NEUTRAL)
        self.assertEqual(self.style(4, 4), OVERVIEW_NEUTRAL)
        self.assertEqual(self.style(4, 4, min_executions=4), OVERVIEW_GREEN)

    def test_overview_cell_defaults_to_the_default_threshold(self):
        _, analysis, _ = inputs([PASSED] * 4, min_executions=4)
        key = analysis.analyses[0]
        self.assertEqual(overview_cell(key).style, OVERVIEW_NEUTRAL)  # 4 < DEFAULT_MIN_EXECUTIONS (5)
        self.assertEqual(overview_cell(key, min_executions=4).style, OVERVIEW_GREEN)
        self.assertEqual(overview_cell(None).value, NO_HISTORY_LABEL)


class OverviewColourTests(unittest.TestCase):
    def assertCell(self, statuses, text, style, **kwargs):
        cell, analysis = overview(statuses, **kwargs)
        self.assertEqual((cell.value, cell.style), (text, style))
        return cell, analysis

    def test_exactly_ninety_percent_is_green(self):
        self.assertCell([PASSED] * 9 + [FAILED], "Passed", OVERVIEW_GREEN)

    def test_just_below_ninety_percent_is_yellow(self):
        self.assertCell([PASSED] * 8 + [FAILED], "Passed", OVERVIEW_YELLOW)  # 8/9 = 88.9%

    def test_exactly_fifty_percent_is_yellow(self):
        self.assertCell([FAILED] + [PASSED] * 5 + [FAILED] * 4, "Failed", OVERVIEW_YELLOW)

    def test_below_fifty_percent_is_red(self):
        self.assertCell([PASSED] * 4 + [FAILED] * 5, "Passed", OVERVIEW_RED)  # 4/9 = 44.4%

    def test_failed_with_ninety_five_percent_is_green(self):
        cell, analysis = self.assertCell([FAILED] + [PASSED] * 19, "Failed", OVERVIEW_GREEN)
        self.assertEqual(analysis.metrics.latest_category, StatusCategory.FAILED)

    def test_passed_with_twenty_percent_is_red(self):
        self.assertCell([PASSED] + [FAILED] * 4, "Passed", OVERVIEW_RED)

    def test_inconsistent_with_ninety_five_percent_is_green(self):
        cell, analysis = self.assertCell([PASSED, FAILED] + [PASSED] * 18, "Passed", OVERVIEW_GREEN)
        self.assertIs(analysis.diagnosis, Diagnosis.INCONSISTENT)
        self.assertIn("Diagnosis: Inconsistent", cell.comment)
        self.assertIn("Pass rate: 95.0%", cell.comment)

    def test_automation_issues_are_conclusive(self):
        # 4 passed / (4 passed + 1 automation issue) = 80%.
        self.assertCell([AUTOMATION_ISSUE] + [PASSED] * 4, "Automation Issue", OVERVIEW_YELLOW)

    def test_minimum_conclusive_executions_is_configurable(self):
        self.assertCell([PASSED] * 5, "Passed", OVERVIEW_GREEN)
        cell, analysis = self.assertCell([PASSED] * 5, "Passed", OVERVIEW_NEUTRAL, min_executions=6)
        self.assertIn("Diagnosis: Insufficient History", cell.comment)
        self.assertCell([PASSED] * 6, "Passed", OVERVIEW_GREEN, min_executions=6)

    def test_blocked_and_inconclusive_executions_do_not_count(self):
        # Five runs, but only four conclusive executions.
        self.assertCell([BLOCKED] + [PASSED] * 4, "Blocked", OVERVIEW_NEUTRAL)
        cell, _ = self.assertCell([BLOCKED] + [PASSED] * 5, "Blocked", OVERVIEW_GREEN)
        self.assertIn("Executions: 6", cell.comment)

    def test_partial_coverage_is_grey(self):
        cell, analysis = self.assertCell([PASSED] * 10, "Partial", OVERVIEW_NEUTRAL, complete=False)
        self.assertIs(analysis.diagnosis, Diagnosis.PARTIAL)
        self.assertIn("Pass rate: 100.0%", cell.comment)

    def test_not_applicable_is_grey(self):
        self.assertCell([NOT_APPLICABLE] + [PASSED] * 6, "Not Applicable", OVERVIEW_NEUTRAL)

    def test_missing_history_is_grey(self):
        collection, analysis, config = inputs([PASSED] * 5)
        sheet = build_report(collection, analysis, config).overview
        headers = [name for name, _ in sheet.columns]
        row = dict(zip(headers, sheet.rows[0], strict=True))
        cell = row["App B Platform X"]
        self.assertEqual((cell.value, cell.style), ("No History", OVERVIEW_NEUTRAL))

    def test_notes_keep_every_detail(self):
        cell, _ = overview([BLOCKED] + [PASSED] * 5)
        lines = [line.split(":", 1)[0] for line in cell.comment.splitlines()]
        self.assertEqual(
            lines, ["Diagnosis", "Confidence", "Executions", "Pass rate", "Coverage", "Latest result", "Blocked reason"]
        )

    def test_details_and_analysis_are_unchanged(self):
        collection, analysis, config = inputs([FAILED] + [PASSED] * 19)
        key = analysis.analyses[0]
        self.assertIs(key.diagnosis, Diagnosis.INCONSISTENT)
        self.assertEqual((key.metrics.pass_rate.numerator, key.metrics.pass_rate.denominator), (19, 20))
        details = build_report(collection, analysis, config).details
        headers = [name for name, _ in details.columns]
        row = dict(zip(headers, details.rows[0], strict=True))
        prefix = "App A | Platform X: "
        self.assertEqual((row[prefix + "Latest Status"].value, row[prefix + "Latest Status"].style), ("Failed", "red"))
        self.assertEqual((row[prefix + "Diagnosis"].value, row[prefix + "Diagnosis"].style), ("Inconsistent", "yellow"))
        self.assertEqual(row[prefix + "Pass Rate"].value, 0.95)


class FormatTests(unittest.TestCase):
    def test_xml_keeps_two_worksheets_without_overview_styles(self):
        collection, analysis, config = inputs([FAILED] + [PASSED] * 19)
        xml = render_spreadsheet(collection, analysis, config)
        self.assertEqual(re.findall(r'<Worksheet ss:Name="([^"]+)"', xml), ["Stability", "Report Info"])
        self.assertNotIn("overview_", xml)
        self.assertIn("Stability Overview background: historical Pass Rate of at least 90%", xml)

    @unittest.skipUnless(openpyxl, "openpyxl is not installed")
    def test_xlsx_backgrounds_follow_the_pass_rate_with_a_neutral_font(self):
        expected = {
            "Failed": ([FAILED] + [PASSED] * 19, "C6EFCE"),
            "Passed": ([PASSED] + [FAILED] * 4, "FFC7CE"),
            "Automation Issue": ([AUTOMATION_ISSUE] + [PASSED] * 4, "FFEB9C"),
            "Blocked": ([BLOCKED] + [PASSED] * 3, "EDEDED"),
        }
        for text, (statuses, fill) in expected.items():
            with self.subTest(text=text):
                collection, analysis, config = inputs(statuses)
                workbook = openpyxl.load_workbook(io.BytesIO(render_xlsx(collection, analysis, config)))
                overview, details = workbook["Stability Overview"], workbook["Stability Details"]
                headers = [cell.value for cell in overview[1]]
                cell = overview.cell(row=2, column=headers.index(COLUMN) + 1)
                self.assertEqual(cell.value, text)
                self.assertEqual(cell.fill.fgColor.rgb[-6:], fill)
                self.assertEqual(cell.font.color.rgb[-6:], "1F1F1F")
                # Details keep the status colours (font and fill) of the latest status.
                headers = [c.value for c in details[1]]
                status = details.cell(row=2, column=headers.index("App A | Platform X: Latest Status") + 1)
                self.assertEqual(status.value, text)
                self.assertNotEqual(status.font.color.rgb[-6:], "1F1F1F")


if __name__ == "__main__":
    unittest.main()
